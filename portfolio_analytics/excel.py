"""Everything in the tearsheet, as an Excel workbook (one sheet per table)."""

from __future__ import annotations

from pathlib import Path

import pandas as pd

from .report import TONE_BACKGROUND
from .verdicts import verdicts_frame

TRADE_IDEA_FIELDS = ("view", "score", "spot", "iv", "iv_source", "rv20", "iv_rv", "vol_regime", "strategy",
                     "rationale", "expiry", "lot_size", "lots", "net_per_lot", "max_loss_per_lot",
                     "breakevens", "futures_alt")


def _sheets(r: dict) -> dict[str, pd.DataFrame]:
    overall, idea, mc = r["overall"], r["index_view"], r["montecarlo"]
    summary_rows = [("Overall verdict", overall["headline"])]
    summary_rows += [(f"Action {i}", action) for i, action in enumerate(overall["actions"], start=1)]
    summary_rows += [(k, round(v, 6) if isinstance(v, float) else v) for k, v in r["summary"].items()]

    trade_idea = pd.concat([
        pd.DataFrame([(k, str(idea[k])) for k in TRADE_IDEA_FIELDS if k in idea], columns=["item", "value"]),
        pd.DataFrame(idea.get("legs", [])),
    ], axis=1)

    return {
        "Summary": pd.DataFrame(summary_rows, columns=["item", "value"]),
        "Verdicts": verdicts_frame(r["section_verdicts"]),
        "Positions verdicts": r["position_verdicts"],
        "Stock signals": r["signals"],
        "NIFTY trade idea": trade_idea,
        "NAV": pd.DataFrame({"nav": r["nav"], "flows": r["ledger"].flows, "cash": r["ledger"].cash,
                             "twr": r["returns"], "benchmark_ret": r["bench_returns"]}),
        "Monthly": r["monthly"],
        "Drawdowns": r["drawdowns"],
        "CAPM": pd.DataFrame(list(r["capm"].items()), columns=["metric", "value"]),
        "VaR": r["var"],
        "Kupiec": pd.DataFrame(r["kupiec"]).T,
        "Open positions": r["positions"],
        "Greeks": r["book"],
        "Risk decomposition": r["euler"],
        "Stress grid": r["stress_grid"],
        "Scenarios": r["scenarios"],
        "Monte Carlo": pd.DataFrame({"percentile": list(mc["terminal_pct"]),
                                     "one_year_return": list(mc["terminal_pct"].values()),
                                     "max_drawdown": list(mc["mdd_pct"].values())}),
        "Sizing": pd.DataFrame([(k, v) for k, v in r["sizing"].items() if k != "risk_parity"],
                               columns=["item", "value"]),
        "Risk parity": r["sizing"].get("risk_parity", pd.DataFrame()),
        "Attribution": r["attr_sleeve"],
        "Monthly P&L": r["attr_monthly"],
        "Trades": r["ledger"].trades,
    }


def to_excel(results: dict, path: str | Path) -> Path:
    from openpyxl.styles import Alignment, Font, PatternFill

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    header_font, header_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="2F5D8A")
    tone_fill = {tone: PatternFill("solid", fgColor=colour.lstrip("#")) for tone, colour in TONE_BACKGROUND.items()}

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, frame in _sheets(results).items():
            frame = pd.DataFrame({"": ["(none)"]}) if frame is None or frame.empty else frame.copy()
            for column in frame.columns:  # Excel has no time zones and hates nanoseconds
                if pd.api.types.is_datetime64_any_dtype(frame[column]):
                    frame[column] = frame[column].dt.date
            if isinstance(frame.index, pd.DatetimeIndex):
                frame.index = frame.index.date
            keep_index = not isinstance(frame.index, pd.RangeIndex)

            sheet_name = name[:31]  # Excel's limit
            frame.to_excel(writer, sheet_name=sheet_name, index=keep_index)
            sheet = writer.sheets[sheet_name]
            for cell in sheet[1]:
                cell.font, cell.fill = header_font, header_fill
            for column in sheet.columns:
                longest = max(len(str(c.value)) if c.value is not None else 0 for c in list(column)[:200])
                sheet.column_dimensions[column[0].column_letter].width = min(max(10, longest + 2), 70)
            sheet.freeze_panes = "B2" if keep_index else "A2"

            if "tone" in frame.columns:  # colour verdict rows green / amber / red
                tone_column = list(frame.columns).index("tone") + int(keep_index)
                for row in sheet.iter_rows(min_row=2):
                    fill = tone_fill.get(row[tone_column].value)
                    if fill:
                        for cell in row:
                            cell.fill = fill
                            cell.alignment = Alignment(wrap_text=True, vertical="top")
    return path
