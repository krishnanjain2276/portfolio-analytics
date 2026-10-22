"""Getting your data in with as little typing as possible.

* ``init_folder`` makes an Excel workbook with the right sheets, dropdowns and instructions.
* ``add_trade`` / ``add_cash`` / ``wizard`` append rows without opening Excel.
* ``fetch_prices`` downloads closes for everything you've traded or are watching.
* ``validate`` checks the data and explains problems in plain English.

Broker imports live in ``tradebook.py``.
"""

from __future__ import annotations

import datetime as dt
from pathlib import Path

import pandas as pd

from . import config
from .charges import estimate_fees
from .data import WORKBOOK, clean_prices, fetch_prices_yfinance, read_table
from .instruments import infer_instrument, normalise_symbol, parse_symbol

TRADE_COLUMNS = ["Date", "Symbol", "Side", "Quantity", "Price", "Fees", "Notes"]
CASH_COLUMNS = ["Date", "Amount", "Notes"]
INSTRUMENT_COLUMNS = ["symbol", "asset_class", "multiplier", "underlying", "strike", "expiry",
                      "option_type", "sector"]

INSTRUCTIONS = [
    ("How to use this workbook", ""),
    ("1. Cashflows", "One row per deposit (+) or withdrawal (-). The first row is your starting capital."),
    ("2. Trades", "One row per trade. Side = BUY or SELL. Quantity = shares for stocks, LOTS for futures/options."),
    ("", "Symbol: an NSE stock (RELIANCE) or NSE F&O symbol (NIFTY26OCTFUT, NIFTY26OCT25000CE, NIFTY26O2725000PE)."),
    ("", "Fees: optional. Leave blank and brokerage, STT, exchange charges, SEBI fee, stamp duty and GST "
         "are estimated."),
    ("3. Watchlist", "Optional. Stocks you don't own but want a BUY / SELL / SHORT verdict on."),
    ("4. Prices", "Optional. `python -m portfolio_analytics fetch-prices --data <folder>` downloads stock and index "
                  "closes (needs `pip install yfinance`). Or paste closes here: Date, then one column per symbol."),
    ("", "Futures and options closes are optional. Missing ones are modelled from NIFTY (futures at fair value, "
         "options with Black-Scholes at the implied vol of your own trade)."),
    ("5. Instruments", "Optional. Only for symbols the engine can't decode, or to override a lot size, "
                       "expiry or sector."),
    ("Run it", "python -m portfolio_analytics run --data <folder>"),
    ("", ""),
    ("Example trades", ""),
    ("2026-01-05 | RELIANCE | BUY | 50 | 1450.5 |", "50 shares"),
    ("2026-01-12 | NIFTY26JANFUT | SELL | 1 | 25900 |", "short 1 lot of the January future"),
    ("2026-01-12 | NIFTY26JAN25500PE | BUY | 2 | 180 |", "2 lots of the 25500 put"),
]


# The Excel template


def init_folder(folder: str | Path, overwrite: bool = False) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / WORKBOOK
    if path.exists() and not overwrite:
        raise FileExistsError(f"{path} already exists (use --force to replace it).")

    workbook = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2F5D8A")

    guide = workbook.active
    guide.title = "Start here"
    for line in INSTRUCTIONS:
        guide.append(list(line))
    guide["A1"].font = Font(bold=True, size=14)
    guide.column_dimensions["A"].width = 52
    guide.column_dimensions["B"].width = 120

    def add_sheet(name, columns, widths):
        sheet = workbook.create_sheet(name)
        sheet.append(columns)
        for cell, width in zip(sheet[1], widths, strict=True):
            cell.font, cell.fill = header_font, header_fill
            cell.alignment = Alignment(horizontal="center")
            sheet.column_dimensions[cell.column_letter].width = width
        sheet.freeze_panes = "A2"
        return sheet

    trades = add_sheet("Trades", TRADE_COLUMNS, [13, 24, 8, 11, 12, 10, 40])
    side_dropdown = DataValidation(type="list", formula1='"BUY,SELL"', allow_blank=True)
    side_dropdown.error, side_dropdown.errorTitle = "Use BUY or SELL", "Side"
    trades.add_data_validation(side_dropdown)
    side_dropdown.add("C2:C5000")
    positive = DataValidation(type="decimal", operator="greaterThan", formula1="0", allow_blank=True)
    positive.error = "Must be a positive number (Side gives the direction)"
    trades.add_data_validation(positive)
    positive.add("D2:E5000")
    for row in range(2, 5001):
        trades.cell(row, 1).number_format = "yyyy-mm-dd"

    cashflows = add_sheet("Cashflows", CASH_COLUMNS, [13, 14, 40])
    for row in range(2, 1001):
        cashflows.cell(row, 1).number_format = "yyyy-mm-dd"

    add_sheet("Watchlist", ["Symbol"], [20])
    add_sheet("Instruments", INSTRUMENT_COLUMNS, [22, 12, 11, 12, 10, 12, 12, 14])
    add_sheet("Prices", ["Date", config.BENCHMARK], [13, 12])
    workbook.save(path)
    return path


# Adding rows


def _first_empty_row(sheet) -> int:
    row = sheet.max_row
    while row > 1 and all(cell.value in (None, "") for cell in sheet[row]):
        row -= 1
    return row + 1


def append_rows(folder: Path, sheet_name: str, rows: list[dict], columns: list[str]) -> str:
    """Add rows to a sheet of the workbook (or to <sheet>.csv if there's no workbook).

    Values are matched to the sheet's own header row by name, so it doesn't matter if you
    have reordered the columns or added your own.
    """
    workbook_path = folder / WORKBOOK
    if not workbook_path.exists():
        csv_path = folder / f"{sheet_name.lower()}.csv"
        new = pd.DataFrame(rows)
        new.columns = [c.lower() for c in new.columns]
        if csv_path.exists():
            existing = pd.read_csv(csv_path)
            new = pd.concat([existing, new[[c for c in new.columns if c in existing.columns] or new.columns]],
                            ignore_index=True)
        new.to_csv(csv_path, index=False)
        return csv_path.name

    from openpyxl import load_workbook

    workbook = load_workbook(workbook_path)
    name = next((s for s in workbook.sheetnames if s.lower() == sheet_name.lower()), None)
    if name is None:
        name = sheet_name.capitalize()
        workbook.create_sheet(name).append(columns)
    sheet = workbook[name]
    header = [str(c.value).strip().lower() if c.value is not None else "" for c in sheet[1]]

    row_number = _first_empty_row(sheet)
    for row in rows:
        values = {key.lower(): value for key, value in row.items()}
        for col_number, heading in enumerate(header, start=1):
            value = values.get(heading)
            if value is None:
                continue
            if isinstance(value, (pd.Timestamp, dt.date)):
                value = pd.Timestamp(value).to_pydatetime()
            cell = sheet.cell(row_number, col_number, value)
            if heading == "date":
                cell.number_format = "yyyy-mm-dd"
        row_number += 1
    workbook.save(workbook_path)
    return f"{workbook_path.name} / {name}"


def add_trade(folder, date, symbol, side, quantity, price, fees=None, notes="") -> tuple[str, float]:
    """Append one trade. Returns where it went and the estimated charges."""
    symbol = normalise_symbol(symbol)
    side = side.strip().upper()
    if side not in ("BUY", "SELL"):
        raise ValueError("side must be BUY or SELL")
    quantity, price = abs(float(quantity)), float(price)
    signed = quantity if side == "BUY" else -quantity
    estimate = estimate_fees(infer_instrument(symbol), signed, price, pd.Timestamp(date))
    row = {"Date": pd.Timestamp(date), "Symbol": symbol, "Side": side, "Quantity": quantity,
           "Price": price, "Fees": fees, "Notes": notes}
    return append_rows(Path(folder), "Trades", [row], TRADE_COLUMNS), estimate


def add_cash(folder, date, amount, notes="") -> str:
    row = {"Date": pd.Timestamp(date), "Amount": float(amount), "Notes": notes}
    return append_rows(Path(folder), "Cashflows", [row], CASH_COLUMNS)


def describe(symbol: str) -> str:
    inst = infer_instrument(symbol)
    expiry = inst.expiry.date() if inst.expiry is not None else None
    if inst.asset_class == "equity":
        return f"stock ({inst.sector})"
    if inst.asset_class == "future":
        return f"future, lot {inst.multiplier:g}, expires {expiry}"
    if inst.asset_class == "option":
        kind = "call" if inst.option_type == "C" else "put"
        return f"{inst.strike:g} {kind}, lot {inst.multiplier:g}, expires {expiry}"
    return "index"


def wizard(folder) -> None:
    print("Add trades (leave the symbol blank to stop). Quantity = shares for stocks, LOTS for F&O.")
    today = dt.date.today().isoformat()
    while True:
        symbol = input("Symbol (e.g. RELIANCE, NIFTY26NOVFUT, NIFTY26NOV25000CE): ").strip()
        if not symbol:
            break
        print(f"  -> {normalise_symbol(symbol)}: {describe(symbol)}")
        date = input(f"Date [{today}]: ").strip() or today
        side = (input("BUY or SELL [BUY]: ").strip() or "BUY").upper()
        quantity = input("Quantity (lots for F&O): ").strip()
        price = input("Price: ").strip()
        fees = input("Fees from the contract note (blank = estimate): ").strip()
        try:
            where, estimate = add_trade(folder, date, symbol, side, quantity, price,
                                        float(fees) if fees else None)
            extra = "" if fees else f" (charges will be estimated at about Rs {estimate:,.2f})"
            print(f"  saved to {where}{extra}")
        except Exception as error:  # a typo shouldn't end the session
            print(f"  not saved: {error}")


# Prices


def fetch_prices(folder, start=None, end=None) -> Path:
    """Download closes for every stock you've traded or are watching, plus NIFTY.

    By default it starts 400 days before your first trade, so the 200-day averages behind
    the verdicts have history to work with from day one.
    """
    folder = Path(folder)
    trades = read_table(folder, "trades")
    watchlist = read_table(folder, "watchlist")
    cashflows = read_table(folder, "cashflows")

    symbols = {config.BENCHMARK}
    first_dates = []
    if trades is not None and len(trades):
        trades.columns = [str(c).strip().lower() for c in trades.columns]
        symbols |= set(trades["symbol"].dropna().map(normalise_symbol))
        first_dates.append(pd.to_datetime(trades["date"]).min())
    if watchlist is not None and len(watchlist):
        symbols |= set(watchlist.iloc[:, 0].dropna().map(normalise_symbol))
    if cashflows is not None and len(cashflows):
        first_dates.append(pd.to_datetime(cashflows.iloc[:, 0]).min())
    symbols = {s for s in symbols if parse_symbol(s) is None}  # Yahoo has no F&O contracts

    if start is None:
        first = min(first_dates) if first_dates else pd.Timestamp.today() - pd.Timedelta(days=730)
        start = first - pd.Timedelta(days=400)

    prices = fetch_prices_yfinance(sorted(symbols), start, end)
    path = folder / "prices.csv"
    if path.exists():
        prices = prices.combine_first(clean_prices(pd.read_csv(path)))
    prices.index.name = "date"
    prices.to_csv(path, float_format="%.2f")
    return path


# Validation


def validate(folder) -> list[tuple[str, str]]:
    """(level, message) pairs, level being OK, INFO, WARN or ERROR."""
    from .data import load_folder
    from .portfolio import run_ledger

    messages = []
    try:
        data = load_folder(folder)
    except Exception as error:
        return [("ERROR", str(error))]

    prices = data.prices
    messages.append(("OK", f"{len(data.trades)} trades, {len(data.cashflows)} cash flows, "
                           f"{prices.shape[1]} price columns from {prices.index[0].date()} to "
                           f"{prices.index[-1].date()}."))
    messages += [("INFO", note) for note in data.notes]

    traded = set(data.trades.symbol)
    no_sector = [s for s in traded if data.instruments[s].asset_class == "equity"
                 and data.instruments[s].sector == "Unclassified"]
    if no_sector:
        messages.append(("INFO", f"No sector for {', '.join(sorted(no_sector))}. Add it in the Instruments "
                                 f"sheet to get sector attribution."))

    try:
        ledger = run_ledger(data.instruments, prices, data.trades, data.cashflows)
    except Exception as error:
        messages.append(("ERROR", str(error)))
        return messages
    messages += [("INFO", note) for note in ledger.notes]

    shorted = [s for s in ledger.units.columns
               if data.instruments[s].asset_class == "equity" and (ledger.units[s] < -1e-9).any()]
    if shorted:
        messages.append(("WARN", f"Short stock positions appear for {', '.join(shorted)}. That's usually a "
                                 f"missing BUY or a SELL entered twice."))
    if (ledger.cash < 0).any():
        first = ledger.cash[ledger.cash < 0].index[0].date()
        messages.append(("WARN", f"Cash goes negative from {first}. A deposit may be missing "
                                 f"(or you trade on margin)."))
    if (ledger.nav <= 0).any():
        messages.append(("ERROR", "NAV reaches zero or below, so returns can't be computed."))

    cutoff = prices.index[-1] - pd.Timedelta(days=10)
    stale = [s for s in sorted(traded)
             if not data.instruments[s].is_derivative and s in prices
             and prices[s].last_valid_index() is not None and prices[s].last_valid_index() < cutoff]
    if stale:
        messages.append(("WARN", f"Prices stop early for {', '.join(stale)}; the last close is carried forward."))

    if not any(level == "ERROR" for level, _ in messages):
        messages.append(("OK", f"Data checks passed. Run: python -m portfolio_analytics run --data {folder}"))
    return messages
